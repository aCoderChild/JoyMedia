import io
from unittest.mock import MagicMock, patch

import frappe
from frappe.tests.utils import FrappeTestCase

from joymedia.services.intake import project as project_intake
from joymedia.tests.support.e2e_setup import _generate_image_bytes


class TestInference(FrappeTestCase):
	def test_reference_roles_are_guessed_from_english_and_vietnamese_names(self):
		role = project_intake.infer_reference_role
		self.assertEqual("Character", role("Gau Shera mascot"))
		# Family terms alone are ambiguous; a user selects their reference role.
		self.assertEqual("General", role("Gấu mẹ"))
		self.assertEqual("Product", role("Shera pack"))
		self.assertEqual("Product", role("Hộp giấy"))
		self.assertEqual("Environment", role("Living room"))
		self.assertEqual("General", role("IMG_0042"))

	def test_a_child_character_is_about_half_its_parents_height(self):
		size = project_intake.infer_relative_size
		for child in ("Gấu con (kid mascot)", "Bé Mai", "Mai, the child", "Bear cub", "Baby panda"):
			self.assertEqual(0.55, size(child), child)
		# "con" before a noun is the Vietnamese classifier ("con gấu" is just "a bear"), not "child".
		for adult in ("Gấu mẹ (mother mascot)", "Mẹ Lan", "Con gấu", "Contest winner", "Concept", ""):
			self.assertEqual(1.0, size(adult), adult)

	def test_format_follows_the_first_aspect_ratio_in_the_text(self):
		preset = project_intake.format_preset
		self.assertEqual("Portrait", preset(" 9:16 (Dọc cho TikTok) & 16:9 (TVC)"))
		self.assertEqual("Landscape", preset("16:9 widescreen"))
		self.assertEqual("Square", preset("1:1"))
		self.assertIsNone(preset("any size"))
		self.assertIsNone(preset(None))


def _workbook_bytes(format_text="9:16 (dọc)", duration="10s"):
	from openpyxl import Workbook

	wb = Workbook()
	ws = wb.active
	ws.title = "Demo 1"
	for row in (
		("Tên sản phẩm", "Khăn giấy Shera"),
		("Mục tiêu", "Quảng cáo"),
		("Thời lượng", duration),
		("Định dạng", format_text),
		("Tone & Mood", "Dễ thương"),
		("USP chính", "Mềm và dai"),
		("Nhân vật", "Gấu mẹ"),
		(),
		(
			"Phân cảnh",
			"Thời lượng",
			"Hình ảnh & Góc máy",
			"Hành động",
			"Lời thoại",
			"Âm thanh",
			"Chữ hiển thị",
		),
		("Hook", "0:00 - 0:05", "Cận cảnh", "Gấu rút giấy", "", "", "x"),
		("Kết", "0:05 - 0:10", "Trung cảnh", "Gấu cười", "", "", "x"),
	):
		ws.append(row)
	buffer = io.BytesIO()
	wb.save(buffer)
	return buffer.getvalue()


class TestCreateFromBrief(FrappeTestCase):
	def setUp(self):
		super().setUp()
		commit = patch.object(frappe.db, "commit")
		commit.start()
		self.addCleanup(commit.stop)
		frappe.set_user("Administrator")

	def tearDown(self):
		frappe.db.rollback()
		super().tearDown()

	def _file(self, name, content):
		return frappe.get_doc(
			{"doctype": "File", "file_name": name, "content": content, "is_private": 1}
		).insert(ignore_permissions=True)

	def test_a_pdf_in_the_references_is_expanded_into_its_pictures(self):
		from PIL import Image

		buffer = io.BytesIO()
		Image.new("RGB", (300, 400), (200, 30, 30)).save(buffer, "PDF")
		pdf = self._file("products.pdf", buffer.getvalue())
		extracted = [{"label": "products p1 box", "page": 1, "role": "Product", "has_alpha": True}]
		with patch.object(project_intake, "add_pdf_images_to_project", return_value=extracted) as add_pdf:
			result = project_intake.create_project_from_brief(
				"Intake PDF", video_idea="A short ad", references=[{"file_name": pdf.name}]
			)

		add_pdf.assert_called_once_with(
			result["project"], pdf.name, pages=None
		)  # no References sheet: every page
		self.assertEqual(extracted, result["references"])

	def _picture(self, name, colour):
		from PIL import Image

		buffer = io.BytesIO()
		Image.new("RGB", (120, 160), colour).save(buffer, "PNG")
		return self._file(name, buffer.getvalue())

	def test_new_videos_of_a_known_product_join_its_project(self):
		pack = self._picture("pack.png", (20, 120, 220))
		first = project_intake.create_project_from_brief(
			"Pack film", references=[{"file_name": pack.name, "label": "pack", "role": "Product"}]
		)["project"]
		again, mascot = self._picture("pack copy.png", (20, 120, 220)), self._picture("bear.png", (240, 240, 230))
		self.assertEqual([first], [row.name for row in project_intake.find_product_project([again.name])])
		self.assertEqual([], project_intake.find_product_project([mascot.name]))

		with patch("joymedia.api.creative.create_product_batch", return_value="CRB-TEST") as batch:
			result = project_intake.create_project_from_brief(
				auto_script=True,
				variants=2,
				into_project=first,
				references=[
					{"file_name": again.name, "label": "pack copy", "role": "Product"},
					{"file_name": mascot.name, "label": "bear", "role": "Character"},
				],
			)
		self.assertEqual(first, result["project"])
		# The project had no character yet; its product pictures stay as they are.
		self.assertEqual([{"label": "bear", "role": "Character"}], result["references"])

		other_bear = self._picture("bear two.png", (200, 150, 90))
		with patch("joymedia.api.creative.create_product_batch", return_value="CRB-TEST"):
			result = project_intake.create_project_from_brief(
				auto_script=True,
				into_project=first,
				references=[{"file_name": other_bear.name, "label": "mascot", "role": "Character"}],
			)
		# A second photo of "the mascot" is never added as a second character.
		self.assertEqual([], result["references"])
		self.assertEqual((first,), batch.call_args.args)
		self.assertEqual(2, batch.call_args.kwargs["count"])

	def test_a_workbook_and_two_images_become_a_project_with_format_and_roles(self):
		workbook = self._file("story.xlsx", _workbook_bytes())
		mascot = self._file("mascot.png", _generate_image_bytes())
		pack = self._file("pack.png", _generate_image_bytes() + b"\x00")
		result = project_intake.create_project_from_brief(
			"Intake test",
			workbook_file=workbook.name,
			references=[
				{"file_name": mascot.name, "label": "Gau Shera mascot"},
				{"file_name": pack.name, "label": "Shera pack", "role": ""},
			],
		)
		project = frappe.get_doc("Media Project", result["project"])

		self.assertEqual("Portrait", project.delivery_preset)
		self.assertEqual((720, 1280), (project.delivery_width, project.delivery_height))
		self.assertEqual(10, project.total_duration_seconds)
		self.assertEqual(2, frappe.db.count("Shot", {"media_project": project.name}))
		self.assertEqual({"Character", "Product"}, {row.reference_role for row in project.selected_media})
		self.assertEqual(
			{"gau_shera_mascot", "shera_pack"}, {row.reference_key for row in project.selected_media}
		)

	def test_a_sheet_project_takes_its_brief_from_the_sheet_and_is_changed_only_by_the_sheet(self):
		from joymedia.api.projects import update_project_brief
		from joymedia.services.planning.revisions import improve_project_video_idea

		workbook = self._file("story.xlsx", _workbook_bytes())
		result = project_intake.create_project_from_brief(
			"Sheet", video_idea="An unrelated drone video", workbook_file=workbook.name
		)
		project = frappe.get_doc("Media Project", result["project"])

		self.assertTrue(project.follows_sheet())
		self.assertIn("Khăn giấy Shera", project.video_idea)
		self.assertNotIn("drone", project.video_idea)
		for change in (
			lambda: project.revise_storyboard("add a close-up"),
			lambda: improve_project_video_idea(project.name),
			lambda: update_project_brief(project.name, video_idea="Something else"),
		):
			self.assertRaises(frappe.ValidationError, change)
		# Saving the same brief back (the studio does before generating) is not a change.
		update_project_brief(project.name, video_idea=project.video_idea)
		# Planned creatively, the project is the AI's to change again.
		project.db_set("planning_mode", "Creative")
		self.assertFalse(frappe.get_doc("Media Project", project.name).follows_sheet())

	def test_a_new_child_character_gets_a_child_size_and_a_size_someone_set_is_kept(self):
		mother = self._file("mother.png", _generate_image_bytes())
		kid = self._file("kid.png", _generate_image_bytes() + b"\x01")
		result = project_intake.create_project_from_brief(
			"Sizes",
			video_idea="A short ad",
			references=[
				{"file_name": mother.name, "label": "Gấu mẹ", "role": "Character"},
				{"file_name": kid.name, "label": "Gấu con", "role": "Character"},
			],
		)
		project = frappe.get_doc("Media Project", result["project"])
		sizes = {row.label: row.relative_size for row in project.selected_media}
		self.assertEqual({"Gấu mẹ": 1.0, "Gấu con": 0.55}, sizes)
		row = next(r for r in project.selected_media if r.label == "Gấu con")
		row.relative_size = 0.8
		project.save(ignore_permissions=True)
		self.assertEqual(0.8, frappe.db.get_value("Project Reference", row.name, "relative_size"))

	def test_readiness_lists_what_is_missing_in_order(self):
		result = project_intake.create_project_from_brief("Empty")
		steps = {step["key"]: step for step in result["readiness"]["steps"]}

		self.assertFalse(result["readiness"]["ready"])
		self.assertEqual("todo", steps["story"]["status"])
		self.assertEqual("todo", steps["approval"]["status"])
		self.assertEqual("story", result["readiness"]["next"])

	def test_description_ratio_and_missing_format_are_not_replaced_by_presets(self):
		result = project_intake.create_project_from_brief(
			"Wide brief", video_idea="A product commercial in 21:9."
		)
		project = frappe.get_doc("Media Project", result["project"])
		self.assertEqual((1680, 720), (project.delivery_width, project.delivery_height))
		empty = project_intake.create_project_from_brief("No shape", video_idea="A product commercial.")
		project = frappe.get_doc("Media Project", empty["project"])
		self.assertFalse(project.delivery_width)
		self.assertEqual(
			"todo", next(s for s in empty["readiness"]["steps"] if s["key"] == "format")["status"]
		)

	def test_custom_workbook_ratio_survives_settings_and_quality_change(self):
		workbook = self._file("custom-shape.xlsx", _workbook_bytes(format_text="4:5 advertising"))
		result = project_intake.create_project_from_brief(workbook_file=workbook.name)
		project = frappe.get_doc("Media Project", result["project"])
		self.assertEqual((720, 900), (project.delivery_width, project.delivery_height))
		project.save_video_settings(
			total_duration_seconds=10, delivery_preset="Portrait", export_quality="Studio 1440p60"
		)
		self.assertEqual((1080, 1350), (project.delivery_width, project.delivery_height))
		project.save_video_settings(total_duration_seconds=10, delivery_preset="Portrait", aspect_ratio="1:1")
		self.assertEqual("Square", project.delivery_preset)
		self.assertEqual(project.delivery_width, project.delivery_height)

	def test_a_bad_format_is_rejected(self):
		with self.assertRaises(frappe.ValidationError):
			project_intake.create_project_from_brief("Bad format", format="sideways")


class TestIntakeIsAtomicAndPrivate(FrappeTestCase):
	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")

	def tearDown(self):
		frappe.db.rollback()
		super().tearDown()

	def test_a_failure_after_the_import_leaves_no_half_built_project(self):
		file_doc = frappe.get_doc(
			{"doctype": "File", "file_name": "w.xlsx", "content": _workbook_bytes(), "is_private": 1}
		).insert(ignore_permissions=True)
		before = frappe.db.count("Media Project")
		with patch(
			"joymedia.services.intake.assets.create_media_asset", side_effect=frappe.ValidationError("boom")
		):
			with self.assertRaises(frappe.ValidationError):
				project_intake.create_project_from_brief(
					"Will fail",
					workbook_file=file_doc.name,
					references=[{"file_name": file_doc.name, "label": "x"}],
				)
		self.assertEqual(before, frappe.db.count("Media Project"))

	def test_a_sheet_missing_must_have_fields_cannot_be_created(self):
		from openpyxl import Workbook

		book = Workbook()
		sheet = book.active
		sheet.title = "Demo 1"
		for row in (
			("Tên sản phẩm", "Khăn giấy"),
			("Mục tiêu", "Quảng cáo"),
			("Thời lượng", "10s"),
			("Phân cảnh", "Hành động"),
			("Hook", "Gấu rút giấy"),
		):
			sheet.append(row)
		buffer = io.BytesIO()
		book.save(buffer)
		frappe.set_user("Administrator")
		workbook = frappe.get_doc(
			{"doctype": "File", "file_name": "m.xlsx", "content": buffer.getvalue(), "is_private": 1}
		).insert(ignore_permissions=True)
		demo = project_intake.preview_workbook(workbook.name)["demos"][0]
		self.assertFalse(demo["can_create"])
		self.assertEqual(["usp", "format", "tone_mood", "characters"], demo["missing_required"])
		with self.assertRaises(frappe.ValidationError):
			project_intake.create_project_from_brief(workbook_file=workbook.name)

	def test_duplicate_detection_never_returns_another_users_asset(self):
		from joymedia.services.intake.assets import _get_duplicate_version

		with patch("joymedia.services.intake.assets.frappe.db.sql", return_value=[]) as sql:
			_get_duplicate_version("hash", "AV", "Image", "Product")
		self.assertIn("ma.owner = %s", sql.call_args.args[0])
		self.assertEqual(frappe.session.user, sql.call_args.args[1][-1])


class TestNamingAndLanguage(FrappeTestCase):
	def tearDown(self):
		frappe.db.rollback()
		super().tearDown()

	def test_a_project_without_a_name_is_named_from_its_workbook(self):
		frappe.set_user("Administrator")
		workbook = frappe.get_doc(
			{"doctype": "File", "file_name": "n.xlsx", "content": _workbook_bytes(), "is_private": 1}
		).insert(ignore_permissions=True)
		with patch.object(frappe.db, "commit"):
			result = project_intake.create_project_from_brief(workbook_file=workbook.name)
		self.assertEqual(
			"Khăn giấy Shera — Demo 1",
			frappe.db.get_value("Media Project", result["project"], "project_name"),
		)

	def test_readiness_follows_the_requested_language(self):
		frappe.set_user("Administrator")
		with patch.object(frappe.db, "commit"):
			project = project_intake.create_project_from_brief("Lang test")["project"]
		vi = {s["key"]: s["label"] for s in project_intake.get_project_readiness(project, lang="vi")["steps"]}
		en = {s["key"]: s["label"] for s in project_intake.get_project_readiness(project, lang="en")["steps"]}
		self.assertEqual("Định dạng video", vi["format"])
		self.assertEqual("Video format", en["format"])


class TestWorkbookPreview(FrappeTestCase):
	def tearDown(self):
		frappe.db.rollback()
		super().tearDown()

	def test_the_preview_reports_format_length_and_scenes_without_creating_anything(self):
		frappe.set_user("Administrator")
		workbook = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": "p.xlsx",
				"content": _workbook_bytes("16:9", "10s"),
				"is_private": 1,
			}
		).insert(ignore_permissions=True)
		before = frappe.db.count("Media Project")
		demo = project_intake.preview_workbook(workbook.name)["demos"][0]

		self.assertEqual("Demo 1", demo["name"])
		self.assertEqual("Landscape", demo["format"])
		self.assertEqual(10, demo["duration_seconds"])
		self.assertEqual(2, demo["scene_count"])
		self.assertEqual(before, frappe.db.count("Media Project"))
		self.assertTrue(demo["can_create"])
		self.assertEqual([], demo["missing_required"])
		self.assertFalse(demo["scenes_drafted"])
		fields = {f["key"]: f for f in demo["fields"]}
		self.assertTrue(fields["product_name"]["required"] and fields["product_name"]["found"])
		self.assertNotIn(
			"asset_requirements", fields
		)  # pictures are uploaded; the "Tài nguyên" note is not asked

	def test_a_non_workbook_is_refused(self):
		frappe.set_user("Administrator")
		image = frappe.get_doc(
			{"doctype": "File", "file_name": "a.png", "content": _generate_image_bytes(), "is_private": 1}
		).insert(ignore_permissions=True)
		with self.assertRaises(frappe.ValidationError):
			project_intake.preview_workbook(image.name)


class TestSchedulerTickWarning(FrappeTestCase):
	def tearDown(self):
		frappe.db.rollback()
		super().tearDown()

	def _steps(self, tick):
		frappe.set_user("Administrator")
		with patch.object(frappe.db, "commit"):
			project = project_intake.create_project_from_brief("Tick")["project"]
		with patch("frappe.utils.scheduler.get_scheduler_tick", return_value=tick):
			return {s["key"]: s for s in project_intake.get_project_readiness(project)["steps"]}

	def test_a_slow_scheduler_tick_is_reported_but_does_not_block_generation(self):
		step = self._steps(240)["scheduler"]
		self.assertEqual("note", step["status"])
		self.assertIn("240", step["detail"])

	def test_a_one_minute_tick_is_fine(self):
		self.assertEqual("done", self._steps(60)["scheduler"]["status"])


class TestPdfImages(FrappeTestCase):
	def test_pictures_are_extracted_with_alpha_and_page_text_and_small_ones_are_skipped(self):
		import tempfile
		from pathlib import Path

		from PIL import Image

		from joymedia.services.intake.project import extract_pdf_images

		with tempfile.TemporaryDirectory() as folder:
			path = Path(folder) / "catalogue.pdf"
			big = Image.new("RGB", (300, 400), (200, 30, 30))
			icon = Image.new("RGB", (40, 40), (0, 0, 0))
			big.save(path, "PDF", save_all=True, append_images=[big, icon])
			found = extract_pdf_images(path)
		# The same picture on two pages is reported once; the icon is skipped.
		self.assertEqual(1, len(found))
		self.assertEqual((300, 400), (found[0]["width"], found[0]["height"]))
		self.assertFalse(found[0]["has_alpha"])
		self.assertTrue(found[0]["png"].startswith(b"\x89PNG"))


class TestStudioSettings(FrappeTestCase):
	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")
		from joymedia.api.projects import create_project

		self.project = create_project("Studio settings", "", "A film.")
		self.addCleanup(frappe.db.rollback)

	def test_the_three_modes_are_saved_validated_and_returned(self):
		from joymedia.api.projects import save_project_video_settings

		with patch.object(frappe.db, "commit"):
			saved = save_project_video_settings(
				self.project.name,
				10,
				"Portrait",
				character_mode="Paste exact art",
				soundtrack_mode="Sound effects only",
				planning_mode="Follow the sheet",
			)
		self.assertEqual(
			("Paste exact art", "Sound effects only", "Follow the sheet"),
			(saved["character_mode"], saved["soundtrack_mode"], saved["planning_mode"]),
		)
		with self.assertRaises(frappe.ValidationError), patch.object(frappe.db, "commit"):
			save_project_video_settings(self.project.name, 10, "Portrait", character_mode="Nonsense")

	def test_a_reference_size_is_validated_and_saved(self):
		from joymedia.api.assets import set_reference_size

		with self.assertRaises(frappe.ValidationError):
			set_reference_size(self.project.name, "AV-NONE", 0.5)
		with self.assertRaises(frappe.ValidationError):
			set_reference_size(self.project.name, "AV-NONE", 9)


class TestSheetPictures(FrappeTestCase):
	def tearDown(self):
		frappe.db.rollback()
		super().tearDown()

	def test_a_picture_on_a_scene_row_becomes_that_scenes_illustration(self):
		import tempfile

		from openpyxl import Workbook
		from openpyxl.drawing.image import Image as SheetImage
		from PIL import Image

		from joymedia.api.storyboard import sheet_row_pictures

		project = frappe.get_doc(
			{"doctype": "Media Project", "project_name": "Pictures", "total_duration_seconds": 10}
		).insert(ignore_permissions=True)
		with tempfile.TemporaryDirectory() as folder:
			Image.new("RGB", (40, 30), "red").save(f"{folder}/outro.png")
			book = Workbook()
			book.active.title = "Demo 1"
			book.active.add_image(SheetImage(f"{folder}/outro.png"), "H19")
			book.save(f"{folder}/sheet.xlsx")
			pictures = sheet_row_pictures(f"{folder}/sheet.xlsx", "Demo 1", project.name)
		self.assertEqual([19], list(pictures))
		saved = frappe.get_doc("File", {"file_url": pictures[19]})
		self.assertEqual(("Media Project", project.name), (saved.attached_to_doctype, saved.attached_to_name))
		self.assertTrue(saved.file_name.endswith(".png"))


class TestStudioRefreshNeverWaitsOnARunLock(FrappeTestCase):
	def test_a_held_run_lock_returns_the_stored_state_instead_of_failing(self):
		from contextlib import contextmanager

		from frappe.utils.file_lock import LockTimeoutError

		from joymedia.api import projects

		@contextmanager
		def held(*args, **kwargs):
			raise LockTimeoutError("held by a worker")
			yield

		project = MagicMock(name="PRJ")
		project.name = "PRJ"
		with (
			patch.object(frappe, "get_doc", return_value=project),
			patch.object(frappe, "get_all", return_value=[frappe._dict(name="RUN-1", status="Running")]),
			patch("frappe.utils.synchronization.filelock", held),
			patch("joymedia.services.generation.orchestrator.refresh_run") as refresh,
			patch.object(projects, "get_project_workspace", return_value={"project": "ok"}),
		):
			self.assertEqual({"project": "ok"}, projects.refresh_project_studio("PRJ"))
		refresh.assert_not_called()


class TestReferencesSheetIntake(FrappeTestCase):
	def setUp(self):
		super().setUp()
		commit = patch.object(frappe.db, "commit")
		commit.start()
		self.addCleanup(commit.stop)
		frappe.set_user("Administrator")

	def tearDown(self):
		frappe.db.rollback()
		super().tearDown()

	def test_uploaded_pictures_take_the_name_role_size_and_rule_the_sheet_gives_them(self):
		from openpyxl import load_workbook

		from joymedia.services.intake.template import build_template

		book = load_workbook(io.BytesIO(build_template("vi")))
		buffer = io.BytesIO()
		book.save(buffer)
		workbook = frappe.get_doc(
			{"doctype": "File", "file_name": "kich_ban.xlsx", "content": buffer.getvalue(), "is_private": 1}
		).insert(ignore_permissions=True)
		mai = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": "be_mai.png",
				"content": _generate_image_bytes(),
				"is_private": 1,
			}
		).insert(ignore_permissions=True)
		result = project_intake.create_project_from_brief(
			"References",
			workbook_file=workbook.name,
			references=[{"file_name": mai.name, "label": "IMG_0042"}],
		)
		project = frappe.get_doc("Media Project", result["project"])
		row = next(r for r in project.selected_media if r.label.startswith("Bé Mai"))
		self.assertEqual(
			("Character", 0.55, "Giữ đúng gương mặt và áo vàng"),
			(row.reference_role, row.relative_size, row.keep_rule),
		)
		readiness = {step["key"]: step for step in result["readiness"]["steps"]}
		# The template also names Mẹ Lan, the Sunny bottle and the logo PDF page, which were not uploaded.
		self.assertIn("Mẹ Lan", readiness["sheet_pictures"]["detail"])
		self.assertNotIn("Bé Mai", readiness["sheet_pictures"]["detail"])
