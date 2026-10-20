import itertools
import tempfile
from pathlib import Path
from typing import ClassVar
from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase

from joymedia.services.intake.template import build_template
from joymedia.services.planning import scene_draft
from joymedia.services.planning.workbook import FIELDS, REQUIRED_FIELDS, missing_required, read_workbook


def _read(lang):
	with tempfile.TemporaryDirectory() as folder:
		path = Path(folder) / "template.xlsx"
		path.write_bytes(build_template(lang))
		return read_workbook(path)


class TestTemplate(FrappeTestCase):
	def test_the_template_is_read_back_complete_in_both_languages(self):
		for lang in ("vi", "en"):
			parsed = _read(lang)
			self.assertEqual(1, len(parsed["demos"]), "the guide sheet must not count as a storyboard")
			demo = parsed["demos"][0]
			self.assertEqual([], missing_required(demo["metadata"]))
			self.assertEqual([], demo["unmapped_labels"])
			self.assertEqual(3, len(demo["scenes"]))
			self.assertTrue(demo["has_scene_table"])

	def test_the_template_marks_exactly_the_required_fields(self):
		self.assertEqual(
			{"product_name", "objective", "usp", "duration", "format", "tone_mood", "characters"},
			set(REQUIRED_FIELDS),
		)
		self.assertTrue(all(key in FIELDS for key in REQUIRED_FIELDS))

	def test_an_empty_required_field_is_reported(self):
		from openpyxl import Workbook

		book = Workbook()
		sheet = book.active
		sheet.title = "Brief"
		sheet.append(["Product name", "Sunny"])
		sheet.append(["Objective", "Ads"])
		sheet.append(["Tone", "Playful"])
		with tempfile.TemporaryDirectory() as folder:
			path = Path(folder) / "brief.xlsx"
			book.save(path)
			demo = read_workbook(path)["demos"][0]
		self.assertEqual(["usp", "duration", "format", "characters"], missing_required(demo["metadata"]))
		self.assertFalse(demo["has_scene_table"])


class TestSceneDraft(FrappeTestCase):
	METADATA: ClassVar[dict] = {
		"product_name": "Sunny",
		"usp": "Cuts grease\nGentle on hands",
		"duration": "30s",
		"characters": "- Mai (main character)\n- Lan",
	}

	def test_a_brief_becomes_hook_one_scene_per_usp_and_outro_covering_the_whole_length(self):
		scenes = scene_draft.draft_scenes(self.METADATA, lang_is_english=True, use_model=False)
		self.assertEqual(["Hook", "USP 1", "USP 2", "Outro"], [s["beat_name"] for s in scenes])
		self.assertEqual(0.0, scenes[0]["start_seconds"])
		self.assertAlmostEqual(30.0, scenes[-1]["end_seconds"], places=1)
		for before, after in itertools.pairwise(scenes):
			self.assertEqual(before["end_seconds"], after["start_seconds"])

	def test_plain_text_uses_only_the_brief(self):
		scenes = scene_draft.draft_scenes(self.METADATA, lang_is_english=True, use_model=False)
		self.assertIn("Mai", scenes[0]["action"])
		self.assertIn("Cuts grease", scenes[1]["action"])

	def test_a_model_answer_with_the_wrong_scene_count_is_ignored(self):
		with patch(
			"joymedia.services.planning.production._ask_planner", return_value={"scenes": [{"action": "x"}]}
		):
			scenes = scene_draft.draft_scenes(self.METADATA, lang_is_english=True)
		self.assertIn("Cuts grease", scenes[1]["action"])

	def test_a_model_answer_that_drifts_off_the_brief_is_ignored(self):
		rows = [
			{
				"action": "Mai flips pancakes",
				"visual_camera": "",
				"voice_over": "",
				"audio_cues": "",
				"on_screen_text": "",
			}
		] * 4
		with patch("joymedia.services.planning.production._ask_planner", return_value={"scenes": rows}):
			scenes = scene_draft.draft_scenes(self.METADATA, lang_is_english=True)
		self.assertNotIn("pancakes", " ".join(s["action"] for s in scenes))

	def test_a_valid_model_answer_is_used(self):
		rows = [
			{
				"action": f"Mai does thing {n} with Sunny, cuts grease, gentle hands",
				"visual_camera": "",
				"voice_over": "",
				"audio_cues": "",
				"on_screen_text": "",
			}
			for n in range(4)
		]
		with patch("joymedia.services.planning.production._ask_planner", return_value={"scenes": rows}):
			scenes = scene_draft.draft_scenes(self.METADATA, lang_is_english=True)
		self.assertEqual("Mai does thing 2 with Sunny, cuts grease, gentle hands", scenes[2]["action"])


class TestTemplateDownload(FrappeTestCase):
	def test_download_returns_an_xlsx(self):
		from joymedia.services.intake.template import download_storyboard_template

		download_storyboard_template("en")
		self.assertEqual("download", frappe.response["type"])
		self.assertTrue(frappe.response["filename"].endswith("-en.xlsx"))
		self.assertTrue(frappe.response["filecontent"].startswith(b"PK"))

	def test_the_studio_links_to_a_whitelisted_download(self):
		source = (
			Path(frappe.get_app_path("joymedia")) / "frontend" / "src" / "components" / "NewProjectDialog.vue"
		).read_text(encoding="utf-8")
		self.assertIn("joymedia.services.intake.template.download_storyboard_template", source)
		method = frappe.get_attr("joymedia.services.intake.template.download_storyboard_template")
		self.assertIn(method, frappe.whitelisted)


class TestClientLayout(FrappeTestCase):
	def test_the_template_has_the_layout_of_the_clients_own_sheets(self):
		import io

		from openpyxl import load_workbook

		sheet = load_workbook(io.BytesIO(build_template("vi"))).worksheets[1]
		labels = {f"A{row}": sheet[f"A{row}"].value for row in (1, 2, 3, 6, 7, 8, 9, 10)}
		self.assertEqual(
			{
				"A1": "Tên sản phẩm",
				"A2": "Mục tiêu",
				"A3": "USP chính",
				"A6": "Thời lượng",
				"A7": "Định dạng",
				"A8": "Tone & Mood",
				"A9": "Nhân vật",
				"A10": None,  # no "Tài nguyên": pictures are uploaded instead
			},
			labels,
		)
		self.assertEqual(
			[
				"Phân cảnh",
				"Thời lượng",
				"Hình ảnh & Góc máy (Visual / Camera)",
				"Hành động & Diễn xuất (Action)",
				"Lời thoại / Voice-over (VO)",
				"Âm thanh (SFX & BGM)",
				"Chữ hiển thị (On-screen Text)",
				"Minh hoạ khung hình",
			],
			[sheet.cell(13, column).value for column in range(1, 9)],
		)
		self.assertTrue(sheet["B4"].value and sheet["B5"].value)  # one USP per row
		self.assertEqual("Storyboard", next(iter(sheet.tables)))

	def test_the_references_brand_and_scene_columns_are_read_back(self):
		from joymedia.services.planning.workbook import parse_storyboard_workbook

		for lang in ("vi", "en"):
			with tempfile.TemporaryDirectory() as folder:
				path = Path(folder) / "template.xlsx"
				path.write_bytes(build_template(lang))
				parsed = parse_storyboard_workbook(path)
			self.assertEqual(1, len(parsed["demos"]), "References and Brand are not storyboards")
			self.assertEqual([], parsed["demos"][0]["unmapped_labels"], lang)
			self.assertEqual(
				["Character", "Character", "Product", "Logo", "Environment", "Prop"],
				[r["role"] for r in parsed["references"]],
			)
			logo = parsed["references"][3]
			self.assertEqual(("catalogue.pdf", 2), (logo["file"], logo["pdf_page"]))
			self.assertEqual("", parsed["references"][4]["file"])  # a place the AI draws
			self.assertEqual(
				{
					"logline",
					"key_message",
					"audience",
					"visual_style",
					"camera_style",
					"claims",
					"avoid",
					"mandatories",
				},
				set(parsed["brand"]),
			)
			self.assertEqual("cinematic", parsed["brand"]["camera_style"])
			self.assertEqual(2, len(parsed["brand"]["claims"].splitlines()))
			last = parsed["demos"][0]["scenes"][2]
			self.assertEqual(2, len(last["cast"]))
			self.assertTrue(last["products"] and last["place"] and last["continuity"])

	def test_the_scene_columns_offer_the_reference_names_and_accept_several(self):
		import io

		from openpyxl import load_workbook

		book = load_workbook(io.BytesIO(build_template("vi")))
		sheet = book["Kịch bản mẫu"]
		self.assertEqual(
			["Nhân vật trong cảnh", "Sản phẩm trong cảnh", "Bối cảnh", "Logic / Liên tục"],
			[sheet.cell(13, c).value for c in range(9, 13)],
		)
		self.assertEqual("A13:Q16", sheet.tables["Storyboard"].ref)
		self.assertEqual("Ảnh mở cảnh", sheet["M13"].value)
		self.assertEqual("Ảnh kết cảnh", sheet["N13"].value)
		lists = [v for v in sheet.data_validations.dataValidation if "Tham chiếu" in (v.formula1 or "")]
		self.assertEqual(5, len(lists))
		self.assertTrue(all(not v.showErrorMessage for v in lists))  # "Bé Mai, Mẹ Lan" is typed, not picked
		self.assertEqual(["Hướng dẫn", "Kịch bản mẫu", "Tham chiếu", "Thương hiệu"], book.sheetnames)

	def test_a_drafted_storyboard_follows_the_brand_sheet(self):
		metadata = {
			"product_name": "Sunny",
			"objective": "Ads",
			"usp": "Sạch dầu",
			"duration": "15s",
			"tone_mood": "vui",
			"characters": "Bé Mai",
		}
		with patch("joymedia.services.planning.production._ask_planner", return_value=None) as ask:
			scene_draft.draft_scenes(
				metadata, brand={"claims": "Sạch dầu với một giọt", "avoid": "Không nói 'diệt 100% vi khuẩn'"}
			)
		prompt = ask.call_args.args[0]
		self.assertIn("APPROVED CLAIMS (use only these): Sạch dầu với một giọt", prompt)
		self.assertIn("MUST NOT SAY OR SHOW: Không nói 'diệt 100% vi khuẩn'", prompt)
