from pathlib import Path
from tempfile import NamedTemporaryFile
from unittest import TestCase

from openpyxl import Workbook

from joymedia.services.storyboard_workbook import parse_storyboard_workbook


def _workbook_path():
	book = Workbook()
	for index, (title, duration) in enumerate((("Demo 1", "35s"), ("Demo 2", "25s"))):
		sheet = book.active if index == 0 else book.create_sheet()
		sheet.title = title
		sheet.append(["Tên sản phẩm", "Khăn giấy Shera 300 tờ 4 lớp"])
		sheet.append(["Thời lượng", duration])
		sheet.append(["Phân cảnh", "Thời lượng", "Hình ảnh", "Hành động", "VO", "SFX", "Chữ"])
		if title == "Demo 1":
			sheet.append(["Hook", "0:00 - 0:04", "Cận cảnh", "Đổ nước ép", "Giấy lau cứ mủn nát", "Cốc đổ", "x"])
		else:
			sheet.append(["Hook", "0:00 - 0:03", "Macro", "Rút giấy", "", "ASMR", "x"])
			sheet.append(["Không bụi", "0:04 - 0:09", "Macro", "Ôm giấy", "Yên tâm", "", "ÊM MỀM"])
			sheet.append(["Thấm hút", "0:10 - 0:15", "Trung cảnh", "Đổ sữa", "4 lớp", "", "4 LỚP"])
			sheet.append(["Outro", "0:15 - 0:22", "Chính diện", "CTA", "Shera", "", "THÙNG 30 GÓI"])
	with NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
		book.save(handle.name)
		return Path(handle.name)


class TestStoryboardWorkbook(TestCase):
	def test_shera_style_workbook_preserves_vietnamese_and_reports_timing_gaps(self):
		path = _workbook_path()
		try:
			parsed = parse_storyboard_workbook(path)
		finally:
			path.unlink(missing_ok=True)
		demo_1, demo_2 = parsed["demos"]
		self.assertEqual("Giấy lau cứ mủn nát", demo_1["scenes"][0]["voice_over"])
		self.assertEqual("", demo_1["scenes"][0]["on_screen_text"])
		self.assertEqual(
			[(3.0, 4.0), (9.0, 10.0), (22.0, 25.0)],
			[(warning["start_seconds"], warning["end_seconds"]) for warning in demo_2["timing_warnings"]],
		)
