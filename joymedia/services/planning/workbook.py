"""Reads a storyboard workbook (.xlsx) into the project's fields and scenes.

A sheet has two parts. The **brief block** (labels in column A, values in column B) holds what the film is about; the
**scene table** (a header row, then one row per scene) holds the storyboard. FIELDS and SCENE_COLUMNS below are the
single description of both: which are required for prompting, what each means, the names it may go by (Vietnamese or
English, any accents, any order) and an example. The parser, the downloadable template, the validation shown to the
client and the studio all read this one schema, so they cannot drift apart.

Reading is deterministic. A label the schema does not know can optionally be mapped by the AI planner
(`label_map`, produced by `intake.project`); the importer never lets a model recover rows or invent scene text.
A sheet without a scene table is valid: the scenes are then drafted from the brief (see `planning.scene_draft`).
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

from joymedia.services.runtime.sampling import sampling

NONE_MARKERS = {"", "x", "n/a", "na", "none", "-"}
# Sheets of the downloadable template that explain it; they are not storyboards.
GUIDE_SHEETS = {"huong dan", "guide", "instructions"}

# key -> required, labels (shown to the client), names it may go by, what to write, example
FIELDS = {
	"product_name": {
		"required": True,
		"label": {"vi": "Tên sản phẩm", "en": "Product name"},
		"names": ("tên sản phẩm", "sản phẩm", "tên sp", "product", "product name"),
		"hint": {"vi": "Tên đầy đủ, kèm quy cách.", "en": "Full name with pack size."},
		"example": {"vi": "Nước rửa chén Sunny hương chanh (500ml)", "en": "Sunny lemon dish soap (500 ml)"},
	},
	"objective": {
		"required": True,
		"label": {"vi": "Mục tiêu", "en": "Objective"},
		"names": ("mục tiêu", "mục đích", "objective", "goal", "campaign goal"),
		"hint": {"vi": "Video dùng để làm gì.", "en": "What the video is for."},
		"example": {"vi": "Quảng cáo sản phẩm (Ads)", "en": "Product advertisement (Ads)"},
	},
	"usp": {
		"required": True,
		"label": {"vi": "USP chính", "en": "Main USPs"},
		"names": (
			"usp",
			"usp chính",
			"usps",
			"main usp",
			"main usps",
			"điểm bán hàng",
			"điểm nổi bật",
			"primary usp",
			"selling points",
		),
		"hint": {"vi": "Mỗi dòng một ý (B3–B5).", "en": "One per row (B3–B5)."},
		"example": {
			"vi": "Làm sạch dầu mỡ chỉ với một giọt\nDịu nhẹ, an toàn cho da tay\nHương chanh tươi mát",
			"en": "Cuts grease with one drop\nGentle on hands\nFresh lemon scent",
		},
	},
	"duration": {
		"required": True,
		"label": {"vi": "Thời lượng", "en": "Duration"},
		"names": ("thời lượng", "độ dài", "duration", "length", "video length"),
		"hint": {"vi": "Tổng thời lượng, vd. 30s (tối đa 60s).", "en": "Total length, e.g. 30s (max 60s)."},
		"example": {"vi": "30s", "en": "30s"},
	},
	"format": {
		"required": True,
		"label": {"vi": "Định dạng", "en": "Format"},
		"names": ("định dạng", "tỉ lệ", "tỷ lệ khung hình", "format", "aspect ratio", "aspect"),
		"hint": {
			"vi": "Tỷ lệ đầu ra, vd. 16:9, 4:5, 1:1 hoặc 9:16.",
			"en": "Output ratio, e.g. 16:9, 4:5, 1:1 or 9:16.",
		},
		"example": {
			"vi": "9:16 (Dọc cho TikTok/Reels/Shorts)",
			"en": "9:16 (vertical for TikTok/Reels/Shorts)",
		},
	},
	"tone_mood": {
		"required": True,
		"label": {"vi": "Tone & Mood", "en": "Tone & Mood"},
		"names": ("tone & mood", "tone and mood", "tone", "mood", "phong cách", "giọng điệu", "cảm xúc"),
		"hint": {"vi": "Cảm giác của video.", "en": "How the video should feel."},
		"example": {
			"vi": "Vui nhộn, tươi sáng, ấm áp gia đình",
			"en": "Playful, bright, warm and family-friendly",
		},
	},
	"characters": {
		"required": True,
		"label": {"vi": "Nhân vật", "en": "Characters"},
		"names": ("nhân vật", "nhân vật chính", "characters", "character", "mascot", "cast"),
		"hint": {
			"vi": "Mỗi dòng một nhân vật; ghi 'nhân vật chính' cho vai chính.",
			"en": "One per line; mark the lead 'main character'.",
		},
		"example": {
			"vi": "- Bé Mai (nhân vật chính)\n- Mẹ Lan",
			"en": "- Mai, the child (main character)\n- Lan, her mother",
		},
	},
	"asset_requirements": {
		"required": False,
		"label": {"vi": "Tài nguyên", "en": "Assets"},
		"names": ("tài nguyên", "assets", "asset requirements", "tài liệu", "hình ảnh tham chiếu"),
		"hint": {
			"vi": "Không cần: ảnh được tải lên ở bước 2.",
			"en": "Not needed: pictures are uploaded in step 2.",
		},
		"example": {
			"vi": "Ảnh chai sản phẩm + logo (tệp PDF của hãng)",
			"en": "Bottle photo + logo (the brand's PDF)",
		},
	},
}

# key -> required, labels, names it may go by, what to write, example. Order is the template's column order.
SCENE_COLUMNS = {
	"beat_name": {
		"required": True,
		"label": {"vi": "Phân cảnh", "en": "Scene"},
		"names": ("phân cảnh", "cảnh", "scene", "scenes", "beat", "shot"),
		"hint": {"vi": "Tên ngắn của cảnh.", "en": "Short scene name."},
		"example": {"vi": "Hook (Tình huống mở đầu)", "en": "Hook (opening situation)"},
	},
	"timing": {
		"required": False,
		"label": {"vi": "Thời lượng", "en": "Timing"},
		"names": ("thời lượng", "thời gian", "timing", "time", "duration"),
		"hint": {"vi": "vd. 0:00 - 0:05. Trống: chia đều.", "en": "e.g. 0:00 - 0:05. Blank: shared evenly."},
		"example": {"vi": "0:00 - 0:05", "en": "0:00 - 0:05"},
	},
	"visual_camera": {
		"required": False,
		"label": {"vi": "Hình ảnh & Góc máy", "en": "Visual & Camera"},
		"names": (
			"hình ảnh & góc máy",
			"hình ảnh",
			"góc máy",
			"góc quay",
			"visual",
			"camera",
			"visual / camera",
		),
		"hint": {"vi": "Khung hình và góc máy.", "en": "Shot and camera move."},
		"example": {
			"vi": "Cận cảnh bồn rửa bếp, ánh nắng buổi sáng",
			"en": "Close-up of a kitchen sink in morning sun",
		},
	},
	"action": {
		"required": True,
		"label": {"vi": "Hành động & Diễn xuất", "en": "Action"},
		"names": (
			"hành động & diễn xuất",
			"hành động",
			"diễn xuất",
			"nội dung",
			"mô tả",
			"action",
			"description",
		),
		"hint": {"vi": "Ai làm gì trong cảnh (quan trọng nhất).", "en": "Who does what (most important)."},
		"example": {
			"vi": "Bé Mai làm đổ dầu ăn ra bếp, hoảng hốt tìm khăn lau",
			"en": "Mai spills cooking oil on the counter and panics, reaching for a cloth",
		},
	},
	"voice_over": {
		"required": False,
		"label": {"vi": "Lời thoại / Voice-over", "en": "Voice-over"},
		"names": (
			"lời thoại",
			"lời thoại / voice-over (vo)",
			"voice-over",
			"voice over",
			"voiceover",
			"vo",
			"thoại",
		),
		"hint": {"vi": "Lời đọc (thu riêng).", "en": "Spoken line (recorded separately)."},
		"example": {"vi": '"Dầu mỡ cứng đầu? Chỉ một giọt!"', "en": '"Stubborn grease? Just one drop!"'},
	},
	"audio_cues": {
		"required": False,
		"label": {"vi": "Âm thanh (SFX & BGM)", "en": "Sound (SFX & BGM)"},
		"names": ("âm thanh", "âm thanh (sfx & bgm)", "sfx", "audio", "bgm", "sound"),
		"hint": {"vi": "SFX: … / BGM: …", "en": "SFX: … / BGM: …"},
		"example": {
			"vi": "SFX: tiếng chai đổ + tiếng nước chảy",
			"en": "SFX: a bottle tipping + running water",
		},
	},
	"on_screen_text": {
		"required": False,
		"label": {"vi": "Chữ hiển thị", "en": "On-screen text"},
		"names": (
			"chữ hiển thị",
			"chữ hiển thị (on-screen text)",
			"on-screen text",
			"on screen text",
			"text",
			"caption",
			"phụ đề",
		),
		"hint": {"vi": "Chữ chèn khi dựng; x nếu không có.", "en": "Text added in editing; x for none."},
		"example": {"vi": "SẠCH DẦU MỠ CHỈ 1 GIỌT", "en": "GREASE GONE IN ONE DROP"},
	},
	"illustration_reference": {
		"required": False,
		"label": {"vi": "Minh hoạ khung hình", "en": "Reference link"},
		"names": ("minh hoạ khung hình", "minh hoạ", "tham chiếu", "illustration", "reference", "link"),
		"hint": {"vi": "Link ảnh/video minh hoạ.", "en": "Link to a sample picture or clip."},
		"example": {"vi": "", "en": ""},
	},
	# Optional: who and what is in the scene, by the names of the References sheet. Left empty, the AI reads it from the
	# Action text. Several names are separated by commas.
	"cast": {
		"required": False,
		"label": {"vi": "Nhân vật trong cảnh", "en": "Characters in the scene"},
		"names": (
			"nhân vật trong cảnh",
			"nhân vật xuất hiện",
			"characters in the scene",
			"characters in scene",
			"cast in scene",
			"cast",
		),
		"hint": {
			"vi": "Tên ở trang Tham chiếu, cách nhau dấu phẩy. Trống: AI đọc từ Hành động.",
			"en": "Names from References, comma-separated. Blank: read from Action.",
		},
		"example": {"vi": "Bé Mai, Mẹ Lan", "en": "Mai, Lan"},
	},
	"products": {
		"required": False,
		"label": {"vi": "Sản phẩm trong cảnh", "en": "Products in the scene"},
		"names": (
			"sản phẩm trong cảnh",
			"sản phẩm xuất hiện",
			"products in the scene",
			"products in scene",
			"product in scene",
		),
		"hint": {
			"vi": "Đúng sản phẩm trong cảnh. Trống: không có.",
			"en": "Exact product shown. Blank: none.",
		},
		"example": {"vi": "Sunny 500ml", "en": "Sunny 500 ml"},
	},
	"place": {
		"required": False,
		"label": {"vi": "Bối cảnh", "en": "Place"},
		"names": ("bối cảnh", "địa điểm", "không gian", "place", "location", "setting"),
		"hint": {"vi": "Cùng tên = cùng bối cảnh.", "en": "Same name = same place."},
		"example": {"vi": "Bếp nhà", "en": "Home kitchen"},
	},
	"continuity": {
		"required": False,
		"label": {"vi": "Logic / Liên tục", "en": "Logic / Continuity"},
		"names": (
			"logic / liên tục",
			"logic / continuity",
			"logic",
			"liên tục",
			"continuity",
			"lưu ý",
			"quy tắc cảnh",
		),
		"hint": {"vi": "Điều phải đúng trong cảnh.", "en": "What must hold in the scene."},
		"example": {
			"vi": "Chai không vỡ; dầu chỉ đổ một vệt nhỏ",
			"en": "The bottle does not break; only a small spill",
		},
	},
	"opening_reference": {
		"required": False,
		"label": {"vi": "Ảnh mở cảnh", "en": "Opening reference"},
		"names": ("ảnh mở cảnh", "opening reference"),
		"hint": {
			"vi": "Tên hoặc key ảnh đã chọn cho khung đầu; trống: tạo ảnh từ mô tả.",
			"en": "Selected reference name or key for the first frame; blank: generate the opening from the scene.",
		},
		"example": {"vi": "Ảnh sản phẩm", "en": "Approved opening"},
	},
	"ending_reference": {
		"required": False,
		"label": {"vi": "Ảnh kết cảnh", "en": "Ending reference"},
		"names": ("ảnh kết cảnh", "ending reference"),
		"hint": {
			"vi": "Tên/key ảnh đã chọn cho khung cuối; tuỳ chọn và cần workflow hỗ trợ.",
			"en": "Selected reference name/key for the final frame; optional and requires workflow support.",
		},
		"example": {"vi": "Ảnh tư thế kết thúc", "en": "Approved ending"},
	},
	"shot_framing": {
		"required": False,
		"label": {"vi": "Cỡ cảnh", "en": "Shot framing"},
		"names": ("cỡ cảnh", "shot framing"),
		"hint": {
			"vi": "Tuỳ chọn: wide, full, medium, medium_close_up, close_up, extreme_close_up.",
			"en": "Optional: wide, full, medium, medium_close_up, close_up, extreme_close_up.",
		},
		"example": {"vi": "", "en": ""},
	},
	"shot_angle": {
		"required": False,
		"label": {"vi": "Góc máy cấu trúc", "en": "Shot angle"},
		"names": ("góc máy cấu trúc", "shot angle"),
		"hint": {
			"vi": "Tuỳ chọn: eye_level, low_angle, high_angle, overhead, oblique, over_shoulder.",
			"en": "Optional: eye_level, low_angle, high_angle, overhead, oblique, over_shoulder.",
		},
		"example": {"vi": "", "en": ""},
	},
	"narrative_stage": {
		"required": False,
		"label": {"vi": "Vai trò cảnh", "en": "Narrative stage"},
		"names": ("vai trò cảnh", "narrative stage"),
		"hint": {
			"vi": "Tuỳ chọn: establish, approach, grip, extract, present, unfold, inspect, detail, feel, hero.",
			"en": "Optional: establish, approach, grip, extract, present, unfold, inspect, detail, feel, hero.",
		},
		"example": {"vi": "", "en": ""},
	},
}

# The References sheet: one row per picture (or place the AI draws), by name. Scenes refer to these names.
REFERENCE_SHEETS = {"tham chieu", "references", "reference", "tai lieu tham chieu"}
REFERENCE_COLUMNS = {
	"name": {
		"required": True,
		"label": {"vi": "Tên", "en": "Name"},
		"width": 22,
		"names": ("tên", "tên gọi", "name"),
		"hint": {"vi": "Tên dùng trong kịch bản.", "en": "Name used in the script."},
	},
	"type": {
		"required": True,
		"label": {"vi": "Loại", "en": "Type"},
		"width": 14,
		"names": ("loại", "vai trò ảnh", "type", "kind"),
		"hint": {
			"vi": "Nhân vật, Người thật, Sản phẩm, Logo, Bối cảnh, Đạo cụ.",
			"en": "Character, Real person, Product, Logo, Place, Prop.",
		},
	},
	"file": {
		"required": False,
		"label": {"vi": "Tệp ảnh / trang PDF", "en": "Picture file / PDF page"},
		"width": 30,
		"names": (
			"tệp ảnh / trang pdf",
			"tệp",
			"tệp ảnh",
			"file",
			"picture file / pdf page",
			"picture",
			"ảnh",
		),
		"hint": {
			"vi": "Tệp sẽ tải lên, hoặc tệp PDF + trang. Trống: AI vẽ.",
			"en": "File you upload, or PDF + page. Blank: AI draws it.",
		},
	},
	"description": {
		"required": False,
		"label": {"vi": "Mô tả", "en": "Description"},
		"width": 40,
		"names": ("mô tả", "description", "mô tả ngắn"),
		"hint": {"vi": "Mô tả ngắn.", "en": "Short description."},
	},
	"family_role": {
		"required": False,
		"label": {"vi": "Vai trò", "en": "Role"},
		"width": 12,
		"names": ("vai trò", "vai", "role", "family role"),
		"hint": {"vi": "mẹ, bố, con, ông, bà…", "en": "mother, father, child…"},
	},
	"height": {
		"required": False,
		"label": {"vi": "Chiều cao tương đối", "en": "Relative height"},
		"width": 14,
		"names": ("chiều cao tương đối", "chiều cao", "relative height", "height", "kích thước"),
		"hint": {"vi": "Nhân vật cao nhất = 1.", "en": "Tallest = 1."},
	},
	"keep_rule": {
		"required": False,
		"label": {"vi": "Quy tắc giữ nguyên", "en": "Must stay as is"},
		"width": 40,
		"names": ("quy tắc giữ nguyên", "quy tắc", "giữ nguyên", "must stay as is", "rules", "keep"),
		"hint": {"vi": "Điều không được thay đổi.", "en": "What must not change."},
	},
	"preservation_mode": {
		"required": False,
		"label": {"vi": "Cách giữ hình", "en": "Preservation mode"},
		"width": 24,
		"names": ("cách giữ hình", "preservation mode"),
		"hint": {
			"vi": "source_pixels: giữ pixel gốc; reference_guided: bám ảnh, cần duyệt hình.",
			"en": "source_pixels keeps original pixels; reference_guided uses appearance conditioning and requires visual review.",
		},
	},
}
REFERENCE_TYPES = {
	"Character": ("nhân vật", "mascot", "linh vật", "character"),
	"Person": ("người thật", "người mẫu", "diễn viên", "real person", "person", "actor"),
	"Product": ("sản phẩm", "product", "bao bì", "packshot"),
	"Logo": ("logo", "thương hiệu", "brand mark"),
	"Environment": ("bối cảnh", "không gian", "địa điểm", "place", "location", "environment", "background"),
	# A recurring object (a cup, a plate): described once, the same in every scene that names it.
	"Prop": ("đạo cụ", "vật dụng", "đồ vật", "prop", "object"),
}
REFERENCE_TYPE_LABELS = {
	"Character": {"vi": "Nhân vật", "en": "Character"},
	"Person": {"vi": "Người thật", "en": "Real person"},
	"Product": {"vi": "Sản phẩm", "en": "Product"},
	"Logo": {"vi": "Logo", "en": "Logo"},
	"Environment": {"vi": "Bối cảnh", "en": "Place"},
	"Prop": {"vi": "Đạo cụ", "en": "Prop"},
}

# The Brand sheet: what the film must say and must not say (a creative brief's message, proof and mandatories).
BRAND_SHEETS = {"thuong hieu", "brand", "brand & claims", "thuong hieu & cong dung"}
BRAND_FIELDS = {
	"logline": {
		"label": {"vi": "Câu chuyện (logline)", "en": "Story (logline)"},
		"names": ("câu chuyện", "câu chuyện (logline)", "logline", "story"),
		"hint": {"vi": "1–2 câu: chuyện gì xảy ra.", "en": "1–2 sentences: what happens."},
		"example": {
			"vi": "Bé Mai làm đổ dầu, Mẹ Lan dùng Sunny lau sạch trong tích tắc.",
			"en": "Mai spills oil; Lan cleans it in seconds with Sunny.",
		},
	},
	"key_message": {
		"label": {"vi": "Thông điệp chính", "en": "Key message"},
		"names": ("thông điệp chính", "thông điệp", "key message", "message"),
		"hint": {"vi": "Một câu người xem phải nhớ.", "en": "One sentence to remember."},
		"example": {
			"vi": "Sunny – sạch dầu mỡ chỉ với một giọt.",
			"en": "Sunny – grease gone with one drop.",
		},
	},
	"audience": {
		"label": {"vi": "Khán giả mục tiêu", "en": "Audience"},
		"names": ("khán giả mục tiêu", "đối tượng", "khán giả", "audience", "target audience"),
		"hint": {"vi": "Video dành cho ai.", "en": "Who it is for."},
		"example": {"vi": "Các bà mẹ 25–40 tuổi ở thành phố", "en": "City mothers aged 25–40"},
	},
	"visual_style": {
		"label": {"vi": "Phong cách hình ảnh", "en": "Visual style"},
		"names": ("phong cách hình ảnh", "phong cách", "visual style", "style", "kiểu hình ảnh"),
		"hint": {"vi": "vd. Hoạt hình 3D, người thật.", "en": "e.g. 3D animation, live action."},
		"example": {"vi": "Hoạt hình 3D", "en": "3D animation"},
	},
	"camera_style": {
		"label": {"vi": "Phong cách máy quay", "en": "Camera style"},
		"names": ("phong cách máy quay", "camera style", "camera grammar"),
		"hint": {
			"vi": "cinematic hoặc restrained; hệ thống chọn theo hành động.",
			"en": "cinematic or restrained; movement follows action.",
		},
		"example": {"vi": "cinematic", "en": "cinematic"},
	},
	"camera_grammar": {
		"label": {"vi": "Ngữ pháp máy quay tham chiếu", "en": "Camera reference grammar"},
		"names": ("ngữ pháp máy quay tham chiếu", "camera reference grammar"),
		"hint": {
			"vi": "JSON: nguồn tham chiếu, nguyên tắc dựng và chuyển động theo vai trò cảnh.",
			"en": "JSON: reference sources, editing notes and moves by narrative stage.",
		},
		"example": {"vi": "", "en": ""},
	},
	"claims": {
		"label": {"vi": "Công dụng đã xác nhận", "en": "Approved claims"},
		"names": ("công dụng đã xác nhận", "công dụng", "approved claims", "claims", "reasons to believe"),
		"hint": {"vi": "Công dụng được phép nói, mỗi dòng một ý.", "en": "Allowed claims, one per row."},
		"example": {
			"vi": "Sạch dầu mỡ với một giọt (kết quả thử nghiệm 2025)\nKhông chứa phosphate",
			"en": "Cuts grease with one drop (2025 test)\nPhosphate-free",
		},
	},
	"avoid": {
		"label": {"vi": "Không được nói / thể hiện", "en": "Must not say or show"},
		"names": (
			"không được nói / thể hiện",
			"không được",
			"tránh",
			"must not say or show",
			"avoid",
			"don'ts",
		),
		"hint": {"vi": "Điều cấm nói hoặc thể hiện.", "en": "What is not allowed."},
		"example": {"vi": "Không nói 'diệt 100% vi khuẩn'", "en": "Never say 'kills 100% of germs'"},
	},
	"mandatories": {
		"label": {"vi": "Bắt buộc có", "en": "Must include"},
		"names": ("bắt buộc có", "bắt buộc", "must include", "mandatories"),
		"hint": {"vi": "Logo, CTA, dòng pháp lý…", "en": "Logo, CTA, legal line…"},
		"example": {
			"vi": "Logo Sunny ở cảnh cuối\nCTA: MUA NGAY",
			"en": "Sunny logo in the last scene\nCTA: BUY NOW",
		},
	},
}

REQUIRED_FIELDS = tuple(key for key, spec in FIELDS.items() if spec["required"])
DEFAULT_SCENE_SECONDS = 5.0


def normalise(text) -> str:
	"""Lower-case, no accents, no punctuation or required/optional marks: how labels are compared."""
	text = re.sub(r"\([^)]*\)|\*", " ", str(text or ""))
	text = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
	text = "".join(c for c in text if not unicodedata.combining(c)).casefold()
	return re.sub(r"[^a-z0-9&]+", " ", text).strip()


_FIELD_LOOKUP = {normalise(name): key for key, spec in FIELDS.items() for name in spec["names"]}
_COLUMN_LOOKUP = {normalise(name): key for key, spec in SCENE_COLUMNS.items() for name in spec["names"]}
_REFERENCE_LOOKUP = {
	normalise(name): key for key, spec in REFERENCE_COLUMNS.items() for name in spec["names"]
}
_BRAND_LOOKUP = {normalise(name): key for key, spec in BRAND_FIELDS.items() for name in spec["names"]}
# "Thời lượng" is the total in the brief block and a scene's timing in the table: contexts differ, the label is shared.


def parse_storyboard_workbook(path: str | Path, label_map: dict | None = None) -> dict:
	"""Parse every sheet that holds a brief or a scene table, without changing source copy or timing."""
	workbook = load_workbook(filename=path, data_only=True, read_only=True)
	demos, references, brand = [], [], {}
	for worksheet in workbook.worksheets:
		title = normalise(worksheet.title)
		if title in GUIDE_SHEETS:
			continue
		rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
		if title in REFERENCE_SHEETS or _reference_header(rows)[1]:
			references.extend(_reference_rows(rows))
			continue
		if title in BRAND_SHEETS:
			brand.update(_brand(rows))
			continue
		demo = _parse_sheet(worksheet, label_map or {}, rows)
		if demo["scenes"] or len(demo["metadata"]) >= 3:
			demos.append(demo)
	if not demos:
		raise ValueError("The workbook has neither a brief (product, objective, USP…) nor a scene table.")
	# The References and Brand sheets belong to every storyboard sheet of the workbook.
	return {"format": "joymedia-storyboard-v3", "demos": demos, "references": references, "brand": brand}


def _reference_header(rows):
	"""(row index, {column index: key}) of a References table: a row naming at least Name, Type and one more column."""
	for index, row in enumerate(rows[:5]):
		columns = {}
		for position, cell in enumerate(row):
			key = _REFERENCE_LOOKUP.get(normalise(cell))
			if key and key not in columns.values():
				columns[position] = key
		if {"name", "type"} <= set(columns.values()) and len(columns) >= 3:
			return index, columns
	return None, {}


def _reference_rows(rows):
	"""[{name, role, file, pdf_page, description, family_role, relative_size, keep_rule}] of the References sheet."""
	index, columns = _reference_header(rows)
	if index is None:
		return []
	references = []
	for row in rows[index + 1 :]:
		values = {key: _text(row, position) for position, key in columns.items()}
		if not values.get("name"):
			continue
		kind = normalise(values.get("type"))
		role = next(
			(
				role
				for role, words in REFERENCE_TYPES.items()
				if any(normalise(word) in kind for word in words)
			),
			"General",
		)
		file_text = _meaningful_text(values.get("file", ""))
		page = (
			re.search(r"(?:trang|page|tr\.?|p\.?)\s*(\d+)", file_text, re.I)
			if re.search(r"\.pdf\b", file_text, re.I)
			else None
		)
		height = re.search(r"\d+(?:[.,]\d+)?", values.get("height", ""))
		size = float(height.group(0).replace(",", ".")) if height else None
		preservation = values.get("preservation_mode", "").strip()
		if preservation and preservation not in {"source_pixels", "reference_guided"}:
			raise ValueError("Preservation mode must be source_pixels or reference_guided.")
		references.append(
			{
				"name": values["name"],
				"role": role,
				"file": re.split(r"\s+(?:trang|page|tr\.?|p\.?)\s*\d+", file_text, flags=re.I)[0].strip()
				if page
				else file_text,
				"pdf_page": int(page.group(1)) if page else None,
				"description": values.get("description", ""),
				"family_role": values.get("family_role", ""),
				"relative_size": (size / 100 if size and size > 2 else size) or None,
				"keep_rule": values.get("keep_rule", ""),
				**({"preservation_mode": preservation} if preservation else {}),
			}
		)
	return references


def _brand(rows):
	"""{field: text} of the Brand sheet (labels in column A, values in B; a value may continue on the rows below)."""
	brand, key = {}, None
	for row in rows:
		label, value = _text(row, 0), _meaningful_text(_text(row, 1))
		if label:
			key = _BRAND_LOOKUP.get(normalise(label))
			if key and value:
				brand[key] = value
		elif key and value:
			brand[key] = f"{brand[key]}\n{value}" if brand.get(key) else value
	return brand


def scene_names(value):
	"""The names a scene cell lists: "Bé Mai, Mẹ Lan" or one per line; "x" or blank lists none."""
	return [
		name.strip(" .")
		for name in re.split(r"[,;\n]+", _meaningful_text(str(value or "").strip()))
		if name.strip(" .")
	]


def read_workbook(path: str | Path) -> dict:
	"""Parse a workbook, letting the AI planner place any label in the brief block that the schema does not know.

	Deterministic first: only labels nothing matched are sent, as a short question whose answer is checked against the
	schema's own field keys. A planner that is down or answers off-schema changes nothing.
	"""
	parsed = parse_storyboard_workbook(path)
	unmapped = sorted({label for demo in parsed["demos"] for label in demo["unmapped_labels"]})
	if not unmapped:
		return parsed
	label_map = _ai_label_map(unmapped)
	return parse_storyboard_workbook(path, label_map) if label_map else parsed


def _ai_label_map(labels: list[str]) -> dict:
	from joymedia.services.planning.production import _ask_planner

	keys = ", ".join(f"{key} ({spec['label']['en']})" for key, spec in FIELDS.items())
	prompt = (
		"A client filled in an advertising brief with these labels. For each label say which field it means, or null "
		f"if it means none of them.\nFIELDS: {keys}\nLABELS: {labels}\n"
		'Return {"map": {"<label>": "<field key or null>"}}'
	)
	try:
		answer = (
			_ask_planner(
				prompt,
				system="You map spreadsheet labels to fields. Return JSON only.",
				sampling_profile="workbook_mapping",
			)
			or {}
		).get("map")
	except Exception:
		return {}
	if not isinstance(answer, dict):
		return {}
	return {label: key for label, key in answer.items() if label in labels and key in FIELDS}


def _parse_sheet(sheet, label_map, rows=None) -> dict:
	rows = rows if rows is not None else [tuple(row) for row in sheet.iter_rows(values_only=True)]
	header_index, columns = _find_scene_header(rows)
	brief_rows = rows[:header_index] if header_index is not None else rows
	metadata, unmapped = _metadata(brief_rows, label_map)
	scenes = _scenes(rows, header_index, columns) if header_index is not None else []
	target_seconds = parse_duration_seconds(metadata.get("duration", ""))
	scenes = _fill_timings(scenes, target_seconds)
	return {
		"name": sheet.title,
		"metadata": metadata,
		"unmapped_labels": unmapped,
		"scenes": scenes,
		"has_scene_table": header_index is not None,
		"timing_warnings": timing_warnings(
			[s for s in scenes if s["timing_source"] == "authored"], target_seconds
		)
		if any(s["timing_source"] == "authored" for s in scenes)
		else [],
	}


def _find_scene_header(rows):
	"""The scene table's header row and {column index: scene key}, or (None, {}) when the sheet has no table."""
	for index, row in enumerate(rows):
		columns = {}
		for position, cell in enumerate(row):
			key = _COLUMN_LOOKUP.get(normalise(cell))
			if key and key not in columns.values():
				columns[position] = key
		if "beat_name" in columns.values() and (
			len(columns) >= 2 or normalise(row[0] if row else "") in ("phan canh", "scene", "scenes")
		):
			# Columns whose header is not recognised keep the template's position (a short header such as "VO").
			order = list(SCENE_COLUMNS)
			for position, cell in enumerate(row):
				if (
					position not in columns
					and str(cell or "").strip()
					and position < len(order)
					and order[position] not in columns.values()
				):
					columns[position] = order[position]
			return index, columns
	return None, {}


def _scenes(rows, header_index, columns):
	scenes = []
	for row in rows[header_index + 1 :]:
		values = {key: _text(row, position) for position, key in columns.items()}
		name = values.get("beat_name", "")
		if not name:
			continue
		timing = values.get("timing", "")
		start, end = parse_time_range(timing)
		length = None
		if start is None:
			length = (
				parse_duration_seconds(timing)
				if timing and re.fullmatch(r"\s*\d+(?:[.,]\d+)?\s*(?:s|sec|giây)?\s*", timing, re.I)
				else None
			)
			if not timing and not values.get("action") and not values.get("visual_camera"):
				continue  # a stray note beneath the table, not a scene
		scenes.append(
			{
				"source_row": len(scenes) + header_index + 2,
				"beat_name": name,
				"source_timing": timing if start is not None else "",
				"start_seconds": start,
				"end_seconds": end,
				"duration_seconds": (end - start) if start is not None else length,
				"timing_source": "authored" if start is not None else "missing",
				"visual_camera": values.get("visual_camera", ""),
				**{
					k: _meaningful_text(values.get(k, ""))
					for k in ("shot_framing", "shot_angle", "narrative_stage")
				},
				"action": values.get("action", ""),
				"voice_over": values.get("voice_over", ""),
				"audio_cues": values.get("audio_cues", ""),
				"on_screen_text": _meaningful_text(values.get("on_screen_text", "")),
				"illustration_reference": _meaningful_text(values.get("illustration_reference", "")),
				"opening_reference": _meaningful_text(values.get("opening_reference", "")),
				"ending_reference": _meaningful_text(values.get("ending_reference", "")),
				"cast": scene_names(values.get("cast")),
				"cast_declared": "cast" in values,
				"products": scene_names(values.get("products")),
				"products_declared": "products" in values,
				"place": _meaningful_text(values.get("place", "")),
				"continuity": _meaningful_text(values.get("continuity", "")),
				# A workbook does not imply visual continuity. Editors explicitly choose
				# Continue only after reviewing the planned composition.
				"continuity_mode": "New Scene",
			}
		)
	return scenes


def _fill_timings(scenes, target_seconds):
	"""Give a scene without a time range one: the length it states, else an even share of what is left of the total."""
	if not scenes or all(s["timing_source"] == "authored" for s in scenes):
		return scenes
	authored = sum(s["duration_seconds"] for s in scenes if s["timing_source"] == "authored")
	missing = [s for s in scenes if s["timing_source"] == "missing"]
	stated = sum(s["duration_seconds"] or 0 for s in missing)
	unstated = [s for s in missing if not s["duration_seconds"]]
	if unstated:
		left = (target_seconds - authored - stated) if target_seconds else None
		share = max(2.0, round(left / len(unstated), 2)) if left and left > 0 else DEFAULT_SCENE_SECONDS
		for scene in unstated:
			scene["duration_seconds"] = share
	# Sequential positions: a scene with its own time range keeps it; the others follow the previous scene.
	cursor = 0.0
	for scene in scenes:
		if scene["timing_source"] == "authored":
			cursor = max(cursor, scene["end_seconds"])
			continue
		scene["start_seconds"], scene["end_seconds"] = cursor, round(cursor + scene["duration_seconds"], 3)
		scene["source_timing"] = _clock(cursor, scene["end_seconds"])
		scene["timing_source"] = "filled"
		cursor = scene["end_seconds"]
	return scenes


def _clock(start, end):
	def fmt(value):
		return f"{int(value // 60)}:{int(value % 60):02d}"

	return f"{fmt(start)} - {fmt(end)}"


def _metadata(rows, label_map):
	"""{field: text} from the brief block, plus the labels nothing could place."""
	metadata, unmapped = {}, []
	continuation_key = None
	for row in rows:
		label = normalise(_text(row, 0))
		raw_label = _text(row, 0)
		value = _text(row, 1)
		if not raw_label and continuation_key and value:
			metadata[continuation_key] = f"{metadata[continuation_key]}\n{value}"
			continue
		key = _FIELD_LOOKUP.get(label) or (label_map or {}).get(raw_label) or None
		if key not in FIELDS:
			continuation_key = None
			if raw_label and value:
				unmapped.append(raw_label)
			continue
		if not value:
			continuation_key = key
			continue
		if key in ("usp", "characters") and metadata.get(key):
			metadata[key] = f"{metadata[key]}\n{value}"
		else:
			metadata[key] = value
		continuation_key = key if key in ("usp", "characters") else None
	return metadata, unmapped


def understood_fields(metadata: dict, lang: str = "vi") -> list[dict]:
	"""Each brief field, whether the sheet gave it and its value, for the client to check before anything is created."""
	lang = "en" if lang == "en" else "vi"
	return [
		{
			"key": key,
			"required": spec["required"],
			"label": spec["label"][lang],
			"hint": spec["hint"][lang],
			"found": bool(metadata.get(key)),
			"value": metadata.get(key, ""),
		}
		for key, spec in FIELDS.items()
		if key
		!= "asset_requirements"  # the pictures themselves are uploaded; a note about them is not needed
	]


def missing_required(metadata: dict) -> list[str]:
	return [key for key in REQUIRED_FIELDS if not metadata.get(key)]


def usp_points(metadata: dict) -> list[str]:
	"""The USPs as separate points, one per line, bullets removed."""
	return [
		re.sub(r"^[\-•*\d.)\s]+", "", line).strip()
		for line in str(metadata.get("usp", "")).splitlines()
		if line.strip()
	]


MAIN_MARKERS = re.compile(r"\s*[-–(]?\s*\b(nhân vật chính|main character|lead|chính)\b\)?", re.I)


def character_names(metadata: dict) -> list[dict]:
	"""[{name, main}] from the Characters field: one per line, 'main character' marks the lead."""
	people = []
	for line in str(metadata.get("characters", "")).splitlines():
		line = re.sub(r"^[\-•*\s]+", "", line).strip()
		if not line:
			continue
		main = bool(MAIN_MARKERS.search(line))
		name = re.sub(r"\s*[-–]\s*$", "", MAIN_MARKERS.sub("", line)).strip(" -–()")
		people.append({"name": name or line, "main": main})
	return people


def parse_duration_seconds(value: str) -> float | None:
	match = re.search(r"(\d+(?:[.,]\d+)?)\s*(?:s|sec|seconds|giây)?", str(value or ""), re.I)
	return float(match.group(1).replace(",", ".")) if match else None


def parse_time_range(value: str) -> tuple[float | None, float | None]:
	parts = re.findall(r"(\d{1,2}):(\d{2})(?:[.,](\d+))?", str(value or ""))
	if len(parts) != 2:
		return None, None
	seconds = [
		int(minutes) * 60 + int(second) + float(f"0.{fraction}" if fraction else 0)
		for minutes, second, fraction in parts
	]
	return (seconds[0], seconds[1]) if seconds[1] > seconds[0] else (None, None)


def timing_warnings(scenes: list[dict], target_seconds: float | None) -> list[dict]:
	warnings = []
	previous_end = 0.0
	for scene in sorted(scenes, key=lambda item: item["start_seconds"]):
		if scene["start_seconds"] > previous_end:
			warnings.append(
				{
					"type": "gap",
					"start_seconds": previous_end,
					"end_seconds": scene["start_seconds"],
					"message": f"No authored scene covers {previous_end:g}s–{scene['start_seconds']:g}s.",
				}
			)
		elif scene["start_seconds"] < previous_end:
			warnings.append(
				{
					"type": "overlap",
					"start_seconds": scene["start_seconds"],
					"end_seconds": previous_end,
					"message": f"Storyboard scenes overlap from {scene['start_seconds']:g}s–{previous_end:g}s.",
				}
			)
		previous_end = max(previous_end, scene["end_seconds"])
	if target_seconds and previous_end < target_seconds:
		warnings.append(
			{
				"type": "gap",
				"start_seconds": previous_end,
				"end_seconds": target_seconds,
				"message": f"No authored scene covers {previous_end:g}s–{target_seconds:g}s.",
			}
		)
	return warnings


def _text(row, index: int) -> str:
	return str(row[index] if len(row) > index and row[index] is not None else "").strip()


def _meaningful_text(value: str) -> str:
	return "" if value.casefold().strip() in NONE_MARKERS else value


BRIEF_KEYS = ("product_name", "objective", "usp", "tone_mood", "characters")


def brief_text(metadata: dict, brand: dict | None = None) -> str:
	"""The sheet's creative brief as the project's video idea, labelled in the sheet's own language.

	Format and length are project settings of their own, so they are left out; USPs and characters stay one per line.
	"""
	vietnamese = not all(str(metadata.get(key) or "").isascii() for key in BRIEF_KEYS)
	lang = "vi" if vietnamese else "en"
	lines = []
	for key in BRIEF_KEYS:
		value = str(metadata.get(key) or "").strip()
		if not value:
			continue
		items = [re.sub(r"^[\-•*\s]+", "", line).strip() for line in value.splitlines() if line.strip()]
		label = FIELDS[key]["label"][lang]
		lines.append(
			f"{label}: {items[0]}"
			if len(items) == 1
			else f"{label}:\n" + "\n".join(f"- {item}" for item in items)
		)
	# The Brand sheet's story and message, when the workbook has one.
	for key in ("logline", "key_message"):
		value = " ".join(str((brand or {}).get(key) or "").split())
		if value:
			lines.append(f"{BRAND_FIELDS[key]['label'][lang]}: {value}")
	return "\n".join(lines)
